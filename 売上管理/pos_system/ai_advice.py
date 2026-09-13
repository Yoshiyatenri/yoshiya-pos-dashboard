"""Claude APIによるAI改善提案生成 と Word(.docx)エクスポート"""
import io

try:
    import anthropic
    _anthropic_ok = True
except ImportError:
    _anthropic_ok = False

try:
    from docx import Document
    _docx_ok = True
except ImportError:
    _docx_ok = False

_SYSTEM_PROMPT = (
    "あなたは駄菓子小売チェーンの経営コンサルタントです。"
    "与えられた実績数値だけを根拠に、日本語で店長向けの改善提案を書いてください。"
    "憶測や一般論だけの助言は避け、具体的な部門名・数値に基づいて指摘してください。"
    "売上の増減については、客数の変化によるものか客単価の変化によるものかを"
    "区別して分析し、集客施策（客数を増やす）と客単価向上施策（一人当たりの"
    "購入額を増やす）のどちらを優先すべきかを明確にしてください"
    "（客数データが「データなし」の場合は、この観点での分析は無理に行わないでください）。"
    "出力は「■現状分析」（3〜4文）と「■改善施策」（箇条書き3〜5個、それぞれ理由付き）の"
    "2つの見出しで構成してください。"
)

_OGATA_TEN = "天理店"  # 郊外大型店（車での来店が前提）。それ以外は店頭に人通りがある店舗

_OGATA_GUIDANCE = (
    "この店舗は郊外の大型店です。客数は主に自動車等での来店による集客に左右され、"
    "目的買いのお客様をいかに集客するかが重要です。客単価は品揃えの広さに左右されます。"
    "集客施策は目的買い需要を喚起する観点、客単価向上施策は品揃え拡充の観点で提案してください。"
)

_SONOTA_GUIDANCE = (
    "この店舗は店頭にある程度の人通りがある立地です。客数は店頭の商品構成など目につく商品に"
    "左右され、いかに店に入ってもらうか（ついで買いのお客様の集客）が重要です。"
    "客単価はついで買いなどを誘発する陳列方法に左右されます。"
    "集客施策は店頭・陳列の工夫で入店を促す観点、客単価向上施策はついで買いを誘発する陳列の観点で提案してください。"
)


def _store_type_guidance(store_label: str) -> str:
    """店舗タイプ（郊外大型店／その他店舗）に応じた分析観点を返す"""
    return _OGATA_GUIDANCE if store_label == _OGATA_TEN else _SONOTA_GUIDANCE


def call_ai_advice(store_label: str, period_label: str, metrics_text: str, api_key: str) -> str:
    """数値サマリーをもとにAI改善提案を生成する。失敗時はエラー文言を返す"""
    if not _anthropic_ok:
        return "⚠️ 生成に失敗しました（anthropicパッケージが未インストールです。`pip install anthropic`を実行してください）"
    if not api_key or "XXXXXXXXXX" in api_key:
        return "⚠️ 生成に失敗しました（Anthropic APIキーが設定されていません。config.jsonのanthropic_api_keyを確認してください）"
    try:
        client = anthropic.Anthropic(api_key=api_key)
        response = client.messages.create(
            model="claude-opus-5",
            max_tokens=16000,
            thinking={"type": "adaptive"},
            system=_SYSTEM_PROMPT + "\n" + _store_type_guidance(store_label),
            messages=[{
                "role": "user",
                "content": f"【{store_label}　{period_label}】\n{metrics_text}",
            }],
        )
        text_blocks = [block.text for block in response.content if block.type == "text"]
        advice = "\n".join(text_blocks).strip()
        if not advice:
            return "⚠️ 生成に失敗しました（応答が空でした）"
        if getattr(response, "stop_reason", None) == "max_tokens":
            advice += "\n\n⚠️ 出力が上限に達したため、途中で切れている可能性があります。"
        return advice
    except anthropic.APIStatusError as e:
        return f"⚠️ 生成に失敗しました（APIエラー: HTTP {e.status_code}）"
    except anthropic.APIConnectionError:
        return "⚠️ 生成に失敗しました（API接続エラー。ネットワークを確認してください）"
    except Exception as e:
        return f"⚠️ 生成に失敗しました（{type(e).__name__}）"


def make_advice_docx(results: dict, period_label: str) -> bytes | None:
    """店舗ごとのAI改善提案をまとめた.docxをバイト列で返す"""
    if not _docx_ok:
        return None
    doc = Document()
    doc.add_heading(f"AI改善提案レポート　{period_label}", level=1)
    for store_label, advice_text in results.items():
        doc.add_heading(store_label, level=2)
        for line in advice_text.split("\n"):
            doc.add_paragraph(line)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
