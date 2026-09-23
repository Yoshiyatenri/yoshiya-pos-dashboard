from io import BytesIO

from openpyxl import load_workbook

import db as db_module


def _login_admin(client):
    client.post("/login", data={"password": "admin123"})


def test_excel_download_returns_xlsx(client):
    _login_admin(client)
    resp = client.get("/admin/goals/results.xlsx")
    assert resp.status_code == 200
    assert resp.mimetype == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def test_excel_contains_all_employees_goal_scores(client):
    _login_admin(client)
    client.post("/admin/employees/add", data={"name": "山田花子"})
    client.post("/admin/employees/add", data={"name": "鈴木一郎"})

    conn = db_module.get_conn()
    cur = conn.cursor()
    cur.execute(db_module.fix("SELECT id FROM employees WHERE name = %s"), ("山田花子",))
    yamada_id = cur.fetchone()[0]
    cur.execute(db_module.fix("SELECT id FROM employees WHERE name = %s"), ("鈴木一郎",))
    suzuki_id = cur.fetchone()[0]
    conn.close()

    client.post("/admin/goals/add", data={"employee_id": str(yamada_id), "goal_text": "接客満足度を上げる"})
    client.post("/admin/goals/add", data={"employee_id": str(suzuki_id), "goal_text": "品出しを早くする"})

    conn = db_module.get_conn()
    cur = conn.cursor()
    cur.execute(db_module.fix("SELECT id FROM goals WHERE employee_id = %s"), (yamada_id,))
    yamada_goal_id = cur.fetchone()[0]
    conn.close()

    client.get("/logout")
    client.post("/login", data={"password": "general123"})
    client.post("/select-employee", data={"employee_id": str(yamada_id)})
    client.post("/goals", data={f"score_{yamada_goal_id}": "8", f"comment_{yamada_goal_id}": "よくできた"})
    client.get("/logout")

    _login_admin(client)
    resp = client.get("/admin/goals/results.xlsx")
    wb = load_workbook(BytesIO(resp.data))
    ws = wb.active
    cell_values = [cell.value for row in ws.iter_rows() for cell in row if cell.value is not None]

    assert "山田花子" in cell_values
    assert "接客満足度を上げる" in cell_values
    assert 8 in cell_values
    assert "よくできた" in cell_values
    # 未採点の目標も一覧に含まれる（評価待ちが分かるように）
    assert "鈴木一郎" in cell_values
    assert "品出しを早くする" in cell_values


def test_excel_requires_admin(client):
    resp = client.get("/admin/goals/results.xlsx", follow_redirects=True)
    assert "パスワードを入力".encode("utf-8") in resp.data
