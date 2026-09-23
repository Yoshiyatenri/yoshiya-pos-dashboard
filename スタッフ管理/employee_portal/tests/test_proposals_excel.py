from io import BytesIO

from openpyxl import load_workbook

import db as db_module


def _login_admin(client):
    client.post("/login", data={"password": "admin123"})


def _add_employee_and_login(client, name="提案太郎"):
    _login_admin(client)
    client.post("/admin/employees/add", data={"name": name})
    conn = db_module.get_conn()
    cur = conn.cursor()
    cur.execute(db_module.fix("SELECT id FROM employees WHERE name = %s"), (name,))
    employee_id = cur.fetchone()[0]
    conn.close()

    client.get("/logout")
    client.post("/login", data={"password": "general123"})
    client.post("/select-employee", data={"employee_id": str(employee_id)})
    return employee_id


def test_excel_download_returns_xlsx(client):
    _login_admin(client)
    resp = client.get("/admin/proposals/results.xlsx")
    assert resp.status_code == 200
    assert resp.mimetype == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def test_excel_contains_named_and_anonymous_proposals(client):
    _add_employee_and_login(client)
    client.post("/proposals", data={"content": "レジ横の陳列を変えたい"})
    client.post("/proposals", data={"content": "匿名の意見です", "is_anonymous": "on"})
    client.get("/logout")

    _login_admin(client)
    resp = client.get("/admin/proposals/results.xlsx")
    wb = load_workbook(BytesIO(resp.data))
    ws = wb.active
    cell_values = [cell.value for row in ws.iter_rows() for cell in row if cell.value is not None]

    assert "レジ横の陳列を変えたい" in cell_values
    assert "提案太郎" in cell_values
    assert "匿名の意見です" in cell_values
    assert "匿名" in cell_values
    assert "未対応" in cell_values


def test_excel_requires_admin(client):
    resp = client.get("/admin/proposals/results.xlsx", follow_redirects=True)
    assert "パスワードを入力".encode("utf-8") in resp.data
